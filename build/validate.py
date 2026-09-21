"""Structural and rendered-text checks for the history-first deliverables."""
from pathlib import Path
import json, re, unicodedata, collections, hashlib
from pptx import Presentation
import fitz
ROOT=Path(__file__).resolve().parents[1]
PPTX=ROOT/'PAP_RAD_HMV_Pathways_v2.pptx'
PDF=ROOT/'PAP_RAD_HMV_Pathways_v2.pdf'
D=json.loads((ROOT/'build/slides.json').read_text())
Q=json.loads((ROOT/'qa/text_boxes.json').read_text())
P=Presentation(PPTX);F=fitz.open(PDF)
md=(ROOT/'PAP_RAD_HMV_Companion_v2.md').read_text()
errors=[];bounds=[];missing=[]
for i,s in enumerate(P.slides):
    if (s._element.get('show','1')=='0') != (i>=26):errors.append(f'Hidden flag: {i+1}')
    for sh in s.shapes:
        if sh.left<-.01*914400 or sh.top<-.01*914400 or sh.left+sh.width>P.slide_width+.01*914400 or sh.top+sh.height>P.slide_height+.01*914400:
            bounds.append({'slide':i+1,'name':sh.name})
    if D[i]['id'] not in s.notes_slide.notes_text_frame.text:errors.append(f'Notes ID: {i+1}')
    if f'<a id="{D[i]["id"].lower()}"></a>' not in md:errors.append(f'Markdown ID: {i+1}')
def tokens(x):
    x=unicodedata.normalize('NFKC',x).lower().replace('\u00ad','')
    return re.findall(r'[a-z0-9]+',x)
for i,page in enumerate(F):
    expected=collections.Counter(tokens('\n'.join(x['text'] for x in Q if x['slide']==i+1)))
    actual=collections.Counter(tokens(page.get_text()))
    diff=expected-actual
    if diff:missing.append({'slide':i+1,'tokens':dict(diff)})
if len(P.slides)!=38 or len(F)!=38:errors.append('Slide/page count')
if sum(s['minutes'] for s in D if s['main'])!=40:errors.append('Nominal time total')
if '(#r01)(#r01)' in md:errors.append('Duplicate reference anchors')
for ref in set(re.findall(r'\]\(#(r\d+)\)',md)):
    if f'<a id="{ref}"></a>' not in md:errors.append('Missing reference '+ref)
# User-requested factual style: no old instructional section headings in the companion.
for term in ['### Live narration','### Teaching objective','### Learner prompt','### Boundaries and common errors','**Ask:**','**Debrief:**']:
    if term in md:errors.append('Instructional heading retained: '+term)
report={'version':'2','reference_date':'2026-09-19','slides':len(P.slides),'main_slides':26,'hidden_appendix_slides':12,'pdf_pages':len(F),'nominal_main_minutes':40,'discussion_minutes':5,'companion_word_count':len(md.split()),'shape_boundary_violations':bounds,'missing_rendered_tokens':missing,'preliminary_layout_height_warnings':json.loads((ROOT/'qa/layout_estimates.json').read_text()),'errors':errors,'visual_review':'Rendered all 38 slides; montage review plus full-size inspection of selected history, physiology, claims, policy, and appendix slides. Final files reviewed after layout corrections.','limitations':['Rendered with LibreOffice in this environment, not Microsoft PowerPoint on the presentation computer.','Historical photograph is attributed; unrestricted online redistribution permission is not established.','Nominal timing is an allocation, not a measured rehearsal.']}
(ROOT/'qa/validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
print(json.dumps(report,ensure_ascii=False,indent=2))
if errors or bounds or missing:raise SystemExit(1)
