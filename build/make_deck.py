"""Build editable history-first slides from the factual content manifest."""
from pathlib import Path
import json, math, re, subprocess
from functools import lru_cache
from PIL import ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement

ROOT=Path(__file__).resolve().parents[1]
D=json.loads((ROOT/'build/slides.json').read_text())
R=json.loads((ROOT/'build/references.json').read_text())
prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
prs.core_properties.title='PAP, RADs, HMVs: Historical Pathways and the Current Landscape'
prs.core_properties.author='Brian W. Locke'
prs.core_properties.subject='Pulmonary fellows lecture; history-first factual revision, version 2'
prs.core_properties.keywords='PAP, CPAP, RAD, HMV, NIV, Medicare, 240.9'
W=13.333333;H=7.5;M=.64;CW=W-2*M
FONT='Calibri';INK='172326';TEXT='253438';MUTED='596A70';ACC='1C6970';PALE='ECF3F3';LIGHT='F4F6F6';RULE='CBD5D8';WHITE='FFFFFF';GRAY='9BA7AB'
QA=[];WARN=[]
FONT_REG=subprocess.check_output(['fc-match','-f','%{file}','Calibri']).decode()
FONT_BOLD=subprocess.check_output(['fc-match','-f','%{file}','Calibri:style=Bold']).decode()
SHORT={'R01':'Locke et al. 2024','R02':'Locke & Brown 2023','R03':'Piper & Yee 2014','R04':'Shah et al. 2023','R05':'Hind et al. 2017','R06':'King 2012','R07':'Wunsch 2023, excerpt','R08':'Lassen 1953','R09':'March of Dimes','R10':'Gusman & Wolfe 2025','R11':'Sullivan et al. 1981','R12':'Sanders & Kern 1990','R13':'71 FR 4518–4525 (2006)','R14':'ONMAP 2021','R15':'HHS OIG 2016','R16':'CMS NCD 240.9','R17':'CMS LCD L33800','R18':'CMS LCD L33718','R19':'CMS A55426','R20':'DME MAC coding guidance','R21':'CMS order requirements','R22':'CGS 2026','R23':'Utah Medicaid','R24':'AASM PAP 2019','R25':'ATS OHS 2019','R26':'ATS COPD NIV 2020','R27':'CHEST NMD 2023','R28':'Ward et al. 2005','R29':'Zheng et al. 2022','R30':'Köhnlein et al. 2014','R31':'Struik et al. 2014','R32':'Murphy et al. 2017','R33':'AASM CSA 2025','R34':'AARC 2012','R35':'Jimenez et al. 2023','R36':'Chung et al. 2022','R37':'Chung et al. 2024','R38':'AASM scoring 2012','R39':'Medicare.gov','R40':'CMS coverage process','R41':'CMS NCD 280.1','R44':'CMS NCD 240.4','R45':'Sullivan et al. 1984','R46':'CMS NCD 240.9 FAQ','R47':'HHS OIG, full report 2016','R48':'CMS 2006','R49':'Georgia Medicaid TEFRA history'}

def rgb(x): return RGBColor.from_string(x)
@lru_cache(None)
def fmetric(size,bold=False):return ImageFont.truetype(FONT_BOLD if bold else FONT_REG,round(size*2))
def wrap_text(txt,width,size,bold=False):
    f=fmetric(size,bold);maxwidth=width*144*.965;out=[]
    for raw in str(txt).split('\n'):
        words=raw.split(' ');current=''
        for word in words:
            test=(current+' '+word).strip()
            if current and f.getlength(test)>maxwidth:
                out.append(current);current=word
            else:current=test
        out.append(current)
    return '\n'.join(out)

def text(s,txt,x,y,w,h,size=25,bold=False,color=TEXT,align='left',valign='top',italic=False,wrap=True):
    txt=wrap_text(txt,w,size,bold) if wrap else str(txt)
    sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=sh.text_frame;tf.clear();tf.word_wrap=True;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    tf.vertical_anchor={'top':MSO_ANCHOR.TOP,'middle':MSO_ANCHOR.MIDDLE,'bottom':MSO_ANCHOR.BOTTOM}[valign]
    for i,ln in enumerate(txt.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=ln
        p.font.name=FONT;p.font.size=Pt(size);p.font.bold=bold;p.font.italic=italic;p.font.color.rgb=rgb(color)
        p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[align]
        p.space_before=Pt(0);p.space_after=Pt(0);p.line_spacing=1.06
    estimate=len(txt.split('\n'))*size/72*1.08
    if estimate>h+.06:WARN.append({'slide':len(prs.slides),'text':txt,'estimated_height':estimate,'height':h})
    QA.append({'slide':len(prs.slides),'text':txt,'x':x,'y':y,'w':w,'h':h,'size':size})
    return sh

def rect(s,x,y,w,h,fill=LIGHT,line_color=None):
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill)
    if line_color:sh.line.color.rgb=rgb(line_color);sh.line.width=Pt(.7)
    else:sh.line.fill.background()
    return sh

def line(s,x1,y1,x2,y2,color=RULE,width=1.2,arrow=False,dash=False):
    sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));sh.line.color.rgb=rgb(color);sh.line.width=Pt(width)
    if arrow:
        e=OxmlElement('a:tailEnd');e.set('type','triangle');sh.line._get_or_add_ln().append(e)
    if dash:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        sh.line.dash_style=MSO_LINE_DASH_STYLE.DASH
    return sh

def poly(s,pts,color=ACC,width=2.4):
    p=[(int(Inches(x)),int(Inches(y))) for x,y in pts]
    b=s.shapes.build_freeform(*p[0]);b.add_line_segments(p[1:],close=False);sh=b.convert_to_shape();sh.fill.background();sh.line.color.rgb=rgb(color);sh.line.width=Pt(width);return sh

def dot(s,x,y,r=.065,color=ACC):
    sh=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x-r),Inches(y-r),Inches(2*r),Inches(2*r));sh.fill.solid();sh.fill.fore_color.rgb=rgb(color);sh.line.fill.background();return sh

def label(s,t,x,y,w,size=12,color=ACC):return text(s,t.upper(),x,y,w,.28,size,True,color)
def bottom(s,t,size=20):
    rect(s,M,6.08,CW,.65,PALE);rect(s,M,6.08,.05,.65,ACC)
    text(s,t,M+.17,6.12,CW-.34,.55,size,False,INK,valign='middle')

def facts(s,items,x,y,w,gap=.95,size=25):
    for i,t in enumerate(items):
        line(s,x,y+i*gap,x+.2,y+i*gap,ACC,2)
        text(s,t,x+.38,y-.1+i*gap,w-.38,gap-.04,size)

def card(s,head,body,x,y,w,h,fs=24,hs=26,fill=LIGHT):
    rect(s,x,y,w,h,fill);rect(s,x,y,w,.045,ACC)
    head_w=wrap_text(head,w-.42,hs,True);n=head_w.count('\n')+1;head_h=max(.46,n*hs/72*1.08+.04)
    text(s,head_w,x+.21,y+.18,w-.42,head_h,hs,True,INK,wrap=False)
    text(s,body,x+.21,y+.30+head_h,w-.42,h-head_h-.47,fs)

def grid_table(s,headers,rows,widths=None,y=1.92,h=3.9,fs=24,hfs=21):
    n=len(headers);widths=widths or [CW/n]*n
    assert abs(sum(widths)-CW)<.03
    headh=.52;rh=(h-headh)/len(rows)
    rect(s,M,y,CW,headh,PALE)
    x=M
    for j,head in enumerate(headers):text(s,head,x+.14,y+.08,widths[j]-.28,headh-.12,hfs,True,INK);x+=widths[j]
    for i,row_ in enumerate(rows):
        yy=y+headh+i*rh
        if i%2==1:rect(s,M,yy,CW,rh,LIGHT)
        line(s,M,yy,W-M,yy,RULE,.6)
        x=M
        for j,cell in enumerate(row_):
            text(s,cell,x+.14,yy+.022,widths[j]-.28,rh-.045,fs,j==0,INK if j==0 else TEXT,valign='middle')
            x+=widths[j]
    line(s,M,y+h,W-M,y+h,RULE,.8)

def columns(s,columns,fs=24):
    width=(CW-.36)/2
    for j,(head,items) in enumerate(columns):
        x=M+j*(width+.36)
        rect(s,x,1.92,width,3.98,LIGHT if j==0 else PALE)
        text(s,head,x+.22,2.12,width-.44,.88,26,True,INK)
        gap=2.83/len(items)
        for i,t in enumerate(items):
            # Spacing, rather than rules, separates wrapped facts.
            text(s,t,x+.22,3.00+i*gap,width-.44,gap-.03,fs)

def base(d):
    s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(WHITE)
    if d['kind']!='title':
        label(s,d['section'],M,.22,9.6,12,MUTED)
        text(s,d['title'],M,.66,CW,1.00,33,True,INK)
        line(s,M,1.61,W-M,1.61,RULE,.8)
        if d['main']:
            parts=['I','II','III','IV'];part=d['section'].split(' · ')[0]
            for j,p in enumerate(parts):rect(s,10.57+j*.53,.29,.39,.07,ACC if p==part else RULE)
    line(s,M,6.91,W-M,6.91,RULE,.7)
    cite=' · '.join(SHORT.get(r,r)+' ['+r+']' for r in d['refs'][:2])
    if len(d['refs'])>2:cite+=' · '+', '.join(d['refs'][2:])
    text(s,cite,M,7.045,10.65,.27,10.2,False,MUTED)
    text(s,('BACKUP · ' if not d['main'] else '')+d['id'],W-M-1.33,7.045,1.33,.25,10.2,False,MUTED,align='right')
    # Complete factual context in notes; no presenter prompts or live-narration scripts.
    note=d['id']+' | '+d['title']+'\n'+d['section']+'\n\n'+re.sub(r'^### ','',d['context'],flags=re.M)
    if d['source_note']:note+='\n\nSOURCE QUALIFICATION\n'+d['source_note']
    note+='\n\nREFERENCES\n'+'\n\n'.join(r+': '+R[r]['citation']+'\n'+R[r]['url'] for r in d['context_refs'])
    s.notes_slide.notes_text_frame.text=note
    if not d['main']:s._element.set('show','0')
    return s

for d in D:
    s=base(d);k=d['kind'];c=d['screen']
    if k=='title':
        label(s,'Pulmonary & Critical Care Fellows',M,.56,10,13)
        text(s,d['title'],M,1.43,CW,1.05,53,True,INK)
        text(s,c['subtitle'],M,2.63,10.2,.85,29)
        line(s,M,3.65,4.3,3.65,ACC,3)
        text(s,'Brian W. Locke',M,4.0,6,.52,25,True)
        paths=c['paths'];bw=(CW-.6)/4
        for i,t in enumerate(paths):
            x=M+i*(bw+.2)
            text(s,str(i+1).zfill(2),x,4.91,bw,.36,16,True,ACC)
            text(s,t,x,5.34,bw,1.0,23,True,INK)
        text(s,'History-first revision · September 19, 2026',M,6.61,10,.22,12,False,MUTED)
    elif k=='history_photo':
        label(s,c['year'],M,1.93,5,16)
        facts(s,c['facts'],M,2.62,5.25,1.17,24)
        # Preserve original aspect ratio; the figure is not a verified Copenhagen image.
        s.shapes.add_picture(str(ROOT/'assets/iron_lung_hind2017.jpeg'),Inches(6.2),Inches(1.96),width=Inches(6.46))
        text(s,c['caption'],6.2,6.58,6.45,.24,10,False,MUTED)
        text(s,'The Autumn Ghost · Hannah Wunsch, 2023',M,6.23,5.45,.35,16,False,MUTED)
    elif k in {'history_timeline','pap_timeline','reg_timeline'}:
        events=c['events'];n=len(events);gap=.22;bw=(CW-(n-1)*gap)/n
        text(s,'Selected milestones · not to scale',M,1.79,CW,.22,11,False,MUTED)
        line(s,M+.25,2.83,W-M-.25,2.83,RULE,1.8)
        for i,(yr,head,body) in enumerate(events):
            x=M+i*(bw+gap)
            text(s,yr,x,2.11,bw,.55,28 if n==3 else 25,True,ACC)
            dot(s,x+.10,2.83)
            text(s,head,x,3.23,bw-.05,1.04,26 if n==3 else 24,True,INK)
            text(s,body,x,4.42,bw-.08,1.38,25 if n==3 else 22.5)
        bottom(s,c['bottom'],20 if n==4 else 21)
    elif k in {'two_columns','rad_paths'}:
        columns(s,c['columns'],24 if d['id']!='S18' else 23)
        if d['id']=='S21':text(s,'Selected entry criteria; detailed conditions in A03',M,1.76,CW,.22,11,False,MUTED)
        bottom(s,c['bottom'],20)
    elif k=='osa_traces':
        facts(s,c['facts'],M,2.22,5.25,1.13,25)
        gx=6.70;gw=5.5
        for yy in [2.6,4.05]:line(s,gx,yy,gx+gw,yy,RULE,.7)
        rect(s,8.55,2.04,1.76,2.64,PALE)
        label(s,'Obstructive apnea',8.54,1.85,2.27,12)
        text(s,'Airflow',gx,2.11,1.63,.35,19,True)
        text(s,'Effort',gx,3.51,1.63,.35,19,True)
        pts=[];eff=[]
        for i in range(241):
            t=i/240;xx=gx+t*gw
            airflow=0 if .337<t<.662 else .28*math.sin(2*math.pi*7*t)
            pts.append((xx,2.87-airflow));eff.append((xx,4.3-.3*math.sin(2*math.pi*7*t)))
        poly(s,pts,ACC,2.4);poly(s,eff,INK,2.0)
        text(s,'Airflow stops; effort continues.',gx,5.06,gw,.8,23,True,ACC)
        text(s,'Conceptual tracings; not patient data.',gx,5.68,gw,.25,12,False,MUTED)
        bottom(s,c['bottom'],22)
    elif k=='cpap':
        for i,(v,l) in enumerate(c['stats']):
            x=M+i*3.00
            text(s,v,x,1.95,2.8,.79,44,True,ACC)
            text(s,l,x,2.82,2.80,.78,21)
        text(s,'Sullivan et al. · Lancet 1981',M,3.64,5.6,.30,17,False,MUTED)
        gx=7.0;gy=4.98;gw=5.25
        line(s,gx,4.0,gx,5.60,RULE,1);line(s,gx,5.6,gx+gw,5.6,RULE,1)
        line(s,gx+.12,4.72,gx+gw-.15,4.72,ACC,3)
        text(s,'Continuous airway pressure',gx,3.86,gw,.6,23,True)
        text(s,'Time',gx+gw-.7,5.64,.8,.3,15,False,MUTED)
        text(s,'No independent IPAP–EPAP difference',gx,5.00,gw,.55,19,False,MUTED)
        facts(s,c['facts'],M,4.24,5.86,.62,20)
        bottom(s,c['bottom'],22)
    elif k=='bilevel':
        text(s,c['facts'][0],M,1.9,8.2,.40,25,True)
        text(s,'PS = IPAP − EPAP',8.4,1.91,4.2,.38,23,True,ACC)
        # Equal patient effort patterns; missing third effort exposes the backup breath.
        for j,mode in enumerate(['Bilevel S','Bilevel S/T']):
            x=M+j*6.25;gx=x+.24;gw=5.41
            label(s,mode,x,2.47,5.75,16)
            text(s,'Pressure',x,2.87,1.65,.28,16,False,MUTED)
            text(s,'Effort',x,4.77,1.65,.28,16,False,MUTED)
            py=4.28;hi=3.24
            line(s,gx,py,gx+gw,py,RULE,.7)
            p=[(gx,py)]
            for i in range(3):
                xx=gx+.63+i*1.56
                if i<2 or j==1:p.extend([(xx,py),(xx,hi),(xx+.61,hi),(xx+.61,py)])
            p.append((gx+gw,py));poly(s,p,ACC,2.8)
            for i in range(3):
                xx=gx+.62+i*1.56
                if i<2:poly(s,[(xx-.10,5.22),(xx,5.22),(xx+.11,5.55),(xx+.31,5.55),(xx+.49,5.22),(xx+.72,5.22)],INK,2.0)
                else:line(s,xx-.1,5.22,xx+.73,5.22,INK,2)
            text(s,'IPAP',gx+.04,3.10,.65,.3,14,False,MUTED)
            text(s,'EPAP',gx+.04,4.36,.65,.3,14,False,MUTED)
            if j==1:
                text(s,'Timed breath',gx+3.36,2.88,1.82,.3,15,True,ACC)
                line(s,gx+4.05,3.16,gx+4.05,3.24,ACC,1.2)
            else:text(s,'No third trigger',gx+3.18,3.72,2.1,.40,17,False,MUTED)
        text(s,'Schematic timing and amplitudes',M,5.76,CW,.22,12,False,MUTED)
        bottom(s,c['bottom'],20)
    elif k in {'three_cards','landscape'}:
        bw=(CW-.48)/3
        for i,(head,body) in enumerate(c['cards']):
            x=M+i*(bw+.24)
            card(s,head,body,x,2.02,bw,2.72,25,26,PALE if i==2 else LIGHT)
        if k=='three_cards':
            text(s,'Increasing respiratory load',M,5.04,5.8,.55,26,True)
            text(s,'Reduced capacity or drive',7.12,5.04,5.45,.55,26,True,ACC)
        else:
            text(s,'Shared modes and patient populations',M,5.01,CW,.58,30,True,ACC)
        bottom(s,c['bottom'],20)
    elif k=='category_grid':
        cells=c['cells'];bw=(CW-.34)/2
        for i,t in enumerate(cells):
            x=M+(i%2)*(bw+.34);y=2.00+(i//2)*1.8
            rect(s,x,y,bw,1.52,PALE if i==0 else LIGHT)
            text(s,t,x+.23,y+.24,bw-.46,1.03,29,True,INK,valign='middle')
        bottom(s,c['bottom'],21)
    elif k in {'platforms','use_table','trial_table','glossary','noncopd','followup'}:
        rr=c['rows'];widths=None;fs=24;hh=3.86;yy=1.94
        if k=='platforms':widths=[3.83,5.92,2.3033];fs=24
        elif k=='use_table':widths=[5.57,6.4833];fs=25
        elif k=='trial_table':widths=[2.8,4.70,4.5533];fs=24
        elif k=='glossary':widths=[2.45,9.6033];fs=23.5
        elif k=='noncopd':widths=[4.15,3.17,4.7333];fs=22.5
        elif k=='followup':
            widths=[3.18,4.34,4.5333];fs=25;hh=2.8
            text(s,c['usage'],M,5.09,CW,.55,29,True,ACC)
            text(s,'At first evaluation: qualifying 30-day period. Later: each paid rental month.',M,5.69,CW,.3,17,False,MUTED)
        grid_table(s,rr[0],rr[1:],widths=widths,y=yy,h=hh,fs=fs)
        bottom(s,c['bottom'],19 if k=='followup' else 20)
    elif k=='co2_overlap':
        facts(s,c['facts'],M,2.20,5.36,1.15,24)
        gx=6.91;gw=5.05;y0=5.53
        line(s,gx,2.35,gx,y0,RULE,1);line(s,gx,y0,gx+gw,y0,RULE,1)
        text(s,'CO₂',gx-.04,1.94,1,.35,20,True)
        text(s,'Sleep',gx+gw-1,5.59,1,.3,16,False,MUTED)
        pts=[(0,.24),(.10,.25),(.14,.65),(.22,.40),(.30,.82),(.38,.61),(.46,1.02),(.54,.84),(.63,1.33),(.72,1.18),(.81,1.65),(.88,1.51),(1,1.95)]
        poly(s,[(gx+a*gw,y0-b*1.35) for a,b in pts],ACC,3)
        text(s,'Accumulation with incomplete\ninter-event compensation',7.35,2.64,4.33,.92,23,True,ACC)
        text(s,'Conceptual course; not patient data.',6.90,5.84,5.3,.24,12,False,MUTED)
        bottom(s,c['bottom'],20)
    elif k=='evidence_cards':
        bw=(CW-.34)/2
        for i,(head,body,result) in enumerate(c['cards']):
            x=M+i*(bw+.34)
            card(s,head,body,x,1.97,bw,2.93,23.5,29,PALE if i==1 else LIGHT)
            line(s,x,5.03,x+bw,5.03,ACC,1.4)
            text(s,result,x,5.20,bw,.71,24,True,ACC)
        bottom(s,c['bottom'],19.5)
    elif k=='claims':
        # Native bars with a zero baseline. This chart shows claims, not clinical overuse.
        gx=1.38;gy=5.44;gw=5.85;gh=3.13
        text(s,'E0464 claims',M,1.88,5.8,.42,23,True)
        for val in [0,100000,200000]:
            yy=gy-val/225000*gh
            line(s,gx,yy,gx+gw,yy,RULE,.7)
            text(s,'0' if not val else f'{val//1000}k',M,yy-.12,.58,.3,13,False,MUTED,align='right')
        for i,(val,yr) in enumerate(zip(c['claims'],c['years'])):
            x=gx+.67+i*2.47;bh=val/225000*gh
            rect(s,x,gy-bh,1.34,max(bh,.04),ACC if i else GRAY)
            text(s,f'{val:,}',x-.22,gy-bh-.42,1.86,.35,23,True,INK,align='center')
            text(s,yr,x-.15,gy+.13,1.65,.35,23,True,INK,align='center')
        for i,(num,what) in enumerate(c['stats']):
            y=2.00+i*1.25
            text(s,num,8.10,y,4.35,.66,36 if i==0 else 29,True,ACC)
            text(s,what,8.10,y+.66,4.30,.62,19.5)
        bottom(s,c['bottom'],19.5)
    elif k=='ncd_intro':
        text(s,c['number'],M,2.10,4.55,1.50,74,True,ACC)
        text(s,'National Coverage\nDetermination',M,3.62,4.65,1.01,29,True)
        text(s,c['subtitle'],M,4.93,4.8,.85,24)
        facts(s,c['facts'],6.42,2.32,6.2,.89,26)
        bottom(s,c['bottom'],20)
    elif k=='hmv_paths':
        rect(s,M,1.91,CW,.84,PALE);text(s,c['baseline'],M+.19,2.07,CW-.38,.55,23,True,ACC,valign='middle')
        text(s,'At least one additional characteristic',M,2.87,CW,.38,20,True,INK)
        for i,t in enumerate(c['criteria']):
            y=3.39+i*.65
            text(s,str(i+1),M,y,.5,.39,24,True,ACC)
            text(s,t,M+.72,y,CW-.72,.51,25)
        bottom(s,c['bottom'],21)
    elif k=='workflow':
        stages=c['stages'];bw=(CW-4*.27)/5
        for i,t in enumerate(stages):
            x=M+i*(bw+.27)
            rect(s,x,1.99,bw,1.04,PALE)
            text(s,t,x+.1,2.10,bw-.2,.80,24,True,INK,align='center',valign='middle')
            if i<4:line(s,x+bw+.025,2.51,x+bw+.245,2.51,ACC,1.6,True)
        for i,(actor,role) in enumerate(c['roles']):
            y=3.55+i*.74
            text(s,actor,M,y,3.02,.50,24,True,INK)
            text(s,role,M+3.2,y,CW-3.2,.61,23.5)
        bottom(s,c['bottom'],21)
    elif k=='appendix_table':
        n=len(c['headers']);widths=[4.4,CW-4.4] if n==2 else [2.65,5.10,CW-7.75]
        if d['id']=='A09':widths=[2.6,4.8,CW-7.4]
        if d['id']=='A11':widths=[2.65,5.15,CW-7.80]
        fs=22 if len(c['rows'])>=5 else 23
        grid_table(s,c['headers'],c['rows'],widths=widths,y=1.91,h=4.09,fs=fs,hfs=20)
        bottom(s,c['bottom'],19)
    elif k=='appendix_columns':
        columns(s,[(c['left_title'],c['left']),(c['right_title'],c['right'])],21.5)
        bottom(s,c['bottom'],19)
    else:raise ValueError('Unimplemented slide kind '+k)

for slide in prs.slides:
    for e in slide._element.xpath('.//a:effectRef'): e.set('idx','0')
out=ROOT/'PAP_RAD_HMV_Pathways_v2.pptx';prs.save(out)
# An all-visible intermediate provides a PDF of all appendix pages as well.
for slide in prs.slides:slide._element.set('show','1')
prs.save(ROOT/'qa/PAP_RAD_HMV_Pathways_v2_all_slides.pptx')
(ROOT/'qa/text_boxes.json').write_text(json.dumps(QA,ensure_ascii=False,indent=2))
(ROOT/'qa/layout_estimates.json').write_text(json.dumps(WARN,ensure_ascii=False,indent=2))
print(f'Wrote {out}; {len(prs.slides)} slides; {len(WARN)} preliminary text-height warnings.')
